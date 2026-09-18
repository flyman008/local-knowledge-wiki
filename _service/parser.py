"""文档解析适配层：多格式分派。

原则（方案§5.2）：
- PDF 先取文字层，图表/扫描页标记需渲染（交 vision 模型），引用页码。
- Word/PPTX 取段落+表格+逐页；必要图片交 vision。
- Excel 保留工作表/单元格/表头/单位，不把整表压成散文。
- 图片交多模态 LLM，保留原图。
- 失败如实记录，不冒充已完成。
"""
from __future__ import annotations

import base64
from dataclasses import dataclass, field
from pathlib import Path

import config as config_mod
import llm as llm_mod
import usage as usage_mod


@dataclass
class ParseResult:
    text: str
    kind: str
    ok: bool
    error: str | None = None
    needs_vision: bool = False
    pages: int = 0
    images: list[str] = field(default_factory=list)  # 原图路径
    tables: list[list[list[str]]] = field(default_factory=list)


def _ext_by_name(filename: str) -> str:
    return Path(filename).suffix.lower()


def parse_file(path: str, cfg: config_mod.Config, vision: llm_mod.LLM | None = None) -> ParseResult:
    ext = _ext_by_name(path)
    p = Path(path)
    if not p.exists():
        return ParseResult("", ext, False, error="文件不存在")

    try:
        if ext in (".pdf",):
            return _parse_pdf(p, vision)
        if ext in (".docx",):
            return _parse_docx(p)
        if ext in (".pptx",):
            return _parse_pptx(p)
        if ext in (".xlsx", ".xls"):
            return _parse_xlsx(p)
        if ext in (".md", ".txt"):
            return ParseResult(p.read_text(encoding="utf-8", errors="replace"), ext, True)
        if ext in (".html", ".htm"):
            return _parse_html(p)
        if ext in (".png", ".jpg", ".jpeg", ".gif", ".webp"):
            return _parse_image(p, cfg, vision)
        if ext in (".mp3", ".mp4", ".wav", ".m4a", ".mov", ".avi"):
            return ParseResult("", ext, False, error="音视频待支持（需转写/抽帧工具）")
        return ParseResult("", ext, False, error=f"不支持的类型 {ext}")
    except Exception as e:  # noqa: BLE001
        return ParseResult("", ext, False, error=f"{type(e).__name__}: {e}")


def _parse_pdf(p: Path, vision: llm_mod.LLM | None = None) -> ParseResult:
    import fitz  # pymupdf

    doc = fitz.open(str(p))
    total_pages = len(doc)
    parts: list[str] = []
    needs_vision = False
    vision_pages: list[int] = []

    for i, page in enumerate(doc):
        txt = page.get_text().strip()
        if not txt:
            # 扫描页/图片页：标记需渲染交 vision
            needs_vision = True
            vision_pages.append(i)
        else:
            parts.append(f"\n--- 第 {i+1} 页 ---\n{txt}")

    # 对扫描页做视觉解析（有 vision 模型时渲染交模型，无则如实标记待识别）
    ocr_failed_pages: list[int] = []
    if vision_pages and vision is not None:
        for i in vision_pages:
            page = doc[i]
            try:
                pix = page.get_pixmap(dpi=150)
                b64 = base64.b64encode(pix.tobytes("png")).decode("ascii")
                res = vision.describe_image(
                    b64, "image/png",
                    prompt=f"这是 PDF 第 {i+1} 页（扫描/图片页，无文本层）。请完整转写这一页的文字内容，"
                           f"包括标题、段落、表格数据、关键数字。无法辨认的字不要猜。",
                )
                parts.append(f"\n--- 第 {i+1} 页（扫描页 OCR） ---\n{res.text}")
                # OCR 用量记账（视觉模型，关联任务由 compile 层统一记文本模型时补全）
                try:
                    usage_mod.record_usage(
                        model_id=res.model,
                        tokens_in=res.tokens_in,
                        tokens_out=res.tokens_out,
                        elapsed_ms=res.elapsed_ms,
                        task_id=None,
                    )
                except Exception:  # noqa: BLE001
                    pass  # 记账失败不影响解析
            except llm_mod.LLMError as e:
                ocr_failed_pages.append(i + 1)
                parts.append(f"\n--- 第 {i+1} 页（扫描页，识别失败：{e.kind}） ---\n")
            except Exception as e:  # noqa: BLE001
                ocr_failed_pages.append(i + 1)
                parts.append(f"\n--- 第 {i+1} 页（扫描页，渲染失败：{type(e).__name__}） ---\n")

    doc.close()
    text = "\n".join(parts)

    if not text.strip():
        return ParseResult("", ".pdf", False, error="无文本层，疑似扫描 PDF", needs_vision=True)

    # OCR 失败处理：全部扫描页失败 = 纯扫描 PDF 未能识别，判失败；部分失败 = 保留待识别状态
    if vision_pages and ocr_failed_pages and len(ocr_failed_pages) == len(vision_pages):
        return ParseResult(
            "", ".pdf", False,
            error=f"扫描页 OCR 全部失败（第 {','.join(map(str, ocr_failed_pages))} 页），待重试",
            needs_vision=True,
        )
    partial_vision = needs_vision and (vision is None or bool(ocr_failed_pages))
    err = None
    if ocr_failed_pages:
        err = f"部分扫描页 OCR 失败（第 {','.join(map(str, ocr_failed_pages))} 页）"
    return ParseResult(text, ".pdf", True, needs_vision=partial_vision, error=err, pages=total_pages)


def _parse_docx(p: Path) -> ParseResult:
    import docx

    d = docx.Document(str(p))
    parts: list[str] = []
    tables: list[list[list[str]]] = []
    for para in d.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for tbl in d.tables:
        rows = [[c.text for c in row.cells] for row in tbl.rows]
        tables.append(rows)
        parts.append("\n[表格]\n" + "\n".join(" | ".join(r) for r in rows))
    return ParseResult("\n".join(parts), ".docx", True, tables=tables)


def _parse_pptx(p: Path) -> ParseResult:
    import pptx

    prs = pptx.Presentation(str(p))
    parts: list[str] = []
    tables: list[list[list[str]]] = []
    slide_no = 0
    for slide in prs.slides:
        slide_no += 1
        parts.append(f"\n--- 第 {slide_no} 页 ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = "\n".join(
                    para.text for para in shape.text_frame.paragraphs if para.text.strip()
                )
                if txt:
                    parts.append(txt)
            if shape.has_table:
                rows = [[c.text for c in row.cells] for row in shape.table.rows]
                tables.append(rows)
                parts.append("[表格]\n" + "\n".join(" | ".join(r) for r in rows))
    return ParseResult("\n".join(parts), ".pptx", True, tables=tables, pages=slide_no)


def _parse_xlsx(p: Path) -> ParseResult:
    import openpyxl

    wb = openpyxl.load_workbook(str(p), data_only=True)  # data_only 取缓存值，非公式
    parts: list[str] = []
    tables: list[list[list[str]]] = []
    for ws in wb.worksheets:
        parts.append(f"\n=== 工作表: {ws.title} ===")
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                rows.append(cells)
                parts.append(" | ".join(cells))
        if rows:
            tables.append(rows)
    return ParseResult("\n".join(parts), ".xlsx", True, tables=tables)


def _parse_html(p: Path) -> ParseResult:
    try:
        import trafilatura

        html = p.read_text(encoding="utf-8", errors="replace")
        text = trafilatura.extract(html, include_comments=False, include_tables=True)
        if text:
            return ParseResult(text, ".html", True)
    except Exception:  # noqa: BLE001
        pass
    # 回退：bs4 抽正文
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(p.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    return ParseResult(soup.get_text("\n", strip=True), ".html", True)


def _parse_image(p: Path, cfg: config_mod.Config, vision: llm_mod.LLM | None) -> ParseResult:
    """图片交多模态 LLM。无 vision 时返回 needs_vision=True，保留原图。"""
    if vision is None:
        return ParseResult("", p.suffix.lower(), False, error="待多模态识别", needs_vision=True, images=[str(p)])
    mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
            ".gif": "image/gif", ".webp": "image/webp"}.get(p.suffix.lower(), "image/png")
    b64 = base64.b64encode(p.read_bytes()).decode("ascii")
    try:
        res = vision.describe_image(b64, mime)
        # 图片识别用量记账（视觉模型）
        try:
            usage_mod.record_usage(
                model_id=res.model,
                tokens_in=res.tokens_in,
                tokens_out=res.tokens_out,
                elapsed_ms=res.elapsed_ms,
                task_id=None,
            )
        except Exception:  # noqa: BLE001
            pass
        return ParseResult(res.text, p.suffix.lower(), True, images=[str(p)])
    except llm_mod.LLMError as e:
        return ParseResult("", p.suffix.lower(), False, error=f"识别失败: {e.kind}", needs_vision=True, images=[str(p)])
